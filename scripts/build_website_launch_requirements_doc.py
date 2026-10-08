from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Selected document preset: compact_reference_guide.
# Named overrides:
# - East Asian body font: Microsoft YaHei (for reliable Chinese rendering).
# - Customer-pack first-page title pattern in the site's own restrained brand palette.
# - Form-table text: 9.5 pt to keep dense, editable matrices usable in portrait layout.

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"

INK = "0B132B"
COBALT = "3347B8"
COBALT_DARK = "22378F"
MUTED = "667085"
LINE = "DCE3F0"
HEADER_FILL = "E8EEF5"
SOFT_BLUE = "F4F8FF"
SOFT_IVORY = "F8F6F1"
GREEN = "149455"
GREEN_FILL = "EAF8F0"
WHITE = "FFFFFF"
PLACEHOLDER = "98A2B3"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Column widths must sum to {CONTENT_WIDTH_DXA}, got {sum(widths_dxa)}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_START_END),
        ("end", CELL_MARGIN_START_END),
    ):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(widths_dxa):
                set_cell_width(cell, widths_dxa[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color: str = LINE, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_run_font(
    run,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    fonts.set(qn("w:cs"), FONT_LATIN)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph_runs(paragraph, *, size: float = 9.5, color: str = INK, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold)


def remove_cell_default_paragraph(cell) -> None:
    for paragraph in list(cell.paragraphs):
        if not paragraph.text and len(cell.paragraphs) > 1:
            cell._tc.remove(paragraph._p)


def add_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = INK,
    size: float = 9.5,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    after: float = 0,
) -> None:
    paragraph = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    remove_cell_default_paragraph(cell)


def set_repeat_title(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, COBALT, 18, 10),
        "Heading 2": (13, COBALT, 14, 7),
        "Heading 3": (12, COBALT_DARK, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_field_run(paragraph, instruction: str) -> None:
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    text_run.append(text)
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    paragraph._p.append(fld_char_begin)
    paragraph._p.append(instr_text)
    paragraph._p.append(fld_char_separate)
    paragraph._p.append(text_run)
    paragraph._p.append(fld_char_end)


def configure_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        left = paragraph.add_run("首版认证｜官网上线资料")
        set_run_font(left, size=9, color=MUTED, bold=True)
        spacer = paragraph.add_run("\t")
        set_run_font(spacer, size=9, color=MUTED)
        right = paragraph.add_run("资料需求清单")
        set_run_font(right, size=9, color=MUTED)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))

        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        prefix = p.add_run("北京首版认证有限公司官网资料准备  ·  ")
        set_run_font(prefix, size=8.5, color=MUTED)
        add_field_run(p, "PAGE")
        for run in p.runs:
            set_run_font(run, size=8.5, color=MUTED)


def add_kicker(doc: Document, text: str, *, color: str = COBALT, after: float = 2) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=10, color=color, bold=True)


def add_title_block(doc: Document) -> None:
    add_kicker(doc, "网站上线准备", color=COBALT, after=3)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("北京首版认证有限公司官网\n上线资料需求清单")
    set_run_font(run, size=27, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.25
    run = subtitle.add_run("用于收集上线前必须确认的主体、服务、资质、法律与内容资料。请由业务、法务和运营共同填写。")
    set_run_font(run, size=11.5, color=MUTED)

    table = doc.add_table(rows=2, cols=4)
    set_table_geometry(table, [1500, 3180, 1500, 3180], indent_dxa=0)
    set_table_borders(table, color=LINE, size="4")
    metadata = [
        ("资料负责人", "请填写"),
        ("目标上线日期", "请填写"),
        ("当前版本", "V1.0"),
        ("最后更新", "____年__月__日"),
    ]
    for i, (label, value) in enumerate(metadata):
        row = i // 2
        col = (i % 2) * 2
        set_cell_shading(table.cell(row, col), HEADER_FILL)
        add_cell_text(table.cell(row, col), label, bold=True, color=COBALT_DARK, size=9.5)
        add_cell_text(table.cell(row, col + 1), value, color=PLACEHOLDER, size=9.5)
    for row in table.rows:
        keep_row_together(row)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)

    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [CONTENT_WIDTH_DXA], indent_dxa=0)
    set_table_borders(callout, color="CFE0FF", size="5")
    set_cell_shading(callout.cell(0, 0), SOFT_BLUE)
    cell = callout.cell(0, 0)
    add_cell_text(cell, "填写规则", bold=True, color=COBALT, size=10.5, after=3)
    rules = [
        "不确定的项目可填写“暂不公开”或“待确认”，请勿猜测。",
        "账号密码、邮箱授权码、接口密钥等敏感信息不要写入本文件。",
        "带“上线前必填”的内容会直接影响网站准确性、合规性或表单可用性。",
        "图片请提供原图，建议宽度不低于 1600px，并标注版权、肖像权和脱敏要求。",
    ]
    for index, text in enumerate(rules, start=1):
        add_cell_text(cell, f"{index}. {text}", color=MUTED, size=9.5, after=1)


def add_phase_banner(doc: Document, title: str, subtitle: str, *, required: bool) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1800, 7560], indent_dxa=0)
    set_table_borders(table, color=GREEN if required else LINE, size="4")
    set_cell_shading(table.cell(0, 0), GREEN_FILL if required else HEADER_FILL)
    set_cell_shading(table.cell(0, 1), SOFT_BLUE if required else SOFT_IVORY)
    add_cell_text(
        table.cell(0, 0),
        "上线前必填" if required else "第二阶段补充",
        bold=True,
        color=GREEN if required else COBALT_DARK,
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_cell_text(table.cell(0, 1), title, bold=True, color=INK, size=12, after=2)
    add_cell_text(table.cell(0, 1), subtitle, color=MUTED, size=9.5)
    keep_row_together(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_section_heading(doc: Document, title: str, *, required: bool = True, note: str | None = None) -> None:
    p = doc.add_paragraph(title, style="Heading 1")
    set_repeat_title(p)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after = Pt(6)
    meta.paragraph_format.keep_with_next = True
    label = meta.add_run("优先级：上线前必填" if required else "优先级：建议补充")
    set_run_font(label, size=9.5, color=GREEN if required else COBALT, bold=True)
    if note:
        divider = meta.add_run("  ｜  ")
        set_run_font(divider, size=9.5, color=LINE)
        text = meta.add_run(note)
        set_run_font(text, size=9.5, color=MUTED)


def add_form_table(
    doc: Document,
    fields: list[tuple[str, str]],
    *,
    multiline_labels: set[str] | None = None,
    widths_dxa: list[int] | None = None,
) -> None:
    widths = widths_dxa or [2700, 6660]
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, widths)
    set_table_borders(table)
    multiline_labels = multiline_labels or set()
    for label, hint in fields:
        row = table.add_row()
        set_cell_shading(row.cells[0], HEADER_FILL)
        set_cell_shading(row.cells[1], WHITE)
        add_cell_text(row.cells[0], label, bold=True, color=COBALT_DARK, size=9.5)
        add_cell_text(row.cells[1], hint or "请填写", color=PLACEHOLDER, size=9.5)
        if label in multiline_labels:
            add_cell_text(row.cells[1], "\n", color=PLACEHOLDER, size=9.5)
        keep_row_together(row)


def add_matrix(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    header_fill: str = HEADER_FILL,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_shading(table.cell(0, index), header_fill)
        add_cell_text(
            table.cell(0, index),
            header,
            bold=True,
            color=COBALT_DARK,
            size=9.2,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            set_cell_shading(row.cells[index], WHITE if len(table.rows) % 2 else SOFT_BLUE)
            add_cell_text(row.cells[index], value, color=INK if index == 0 else PLACEHOLDER, size=9)
        keep_row_together(row)


def add_notes_box(doc: Document, prompt: str = "补充说明 / 待确认事项") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=LINE, size="4")
    set_cell_shading(table.cell(0, 0), SOFT_IVORY)
    add_cell_text(table.cell(0, 0), prompt, bold=True, color=COBALT_DARK, size=9.5, after=2)
    add_cell_text(table.cell(0, 0), "请填写：\n\n", color=PLACEHOLDER, size=9.5)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def build_document(output_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    configure_header_footer(doc)

    core_properties = doc.core_properties
    core_properties.title = "北京首版认证有限公司官网上线资料需求清单"
    core_properties.subject = "官网上线资料收集与确认"
    core_properties.author = "北京首版认证有限公司"
    core_properties.keywords = "官网上线, 资料需求, 资质, 服务, 合规"

    add_title_block(doc)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_phase_banner(
        doc,
        "第一阶段：上线前必须确认",
        "以下八类资料直接影响公司主体准确性、服务表达、资质归属、法律合规和联系方式。",
        required=True,
    )
    add_matrix(
        doc,
        ["资料模块", "建议负责人", "当前状态", "预计完成时间"],
        [
            ["企业主体信息", "法务 / 行政", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["联系方式", "运营 / 商务", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["备案与法律信息", "法务 / 技术", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["公司定位与核心文案", "负责人 / 品牌", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["服务与产品信息", "业务负责人", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["星眸 AIPR", "产品 / 技术", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["资质与证书", "法务 / 行政", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
            ["合作伙伴", "商务 / 法务", "[ ] 未开始  [ ] 进行中  [ ] 已确认", "请填写"],
        ],
        [2200, 2100, 3260, 1800],
    )

    add_page_break(doc)
    add_phase_banner(
        doc,
        "第一阶段：上线前必须确认",
        "填写时优先确认对外展示口径和资料归属，所有不确定内容应标记“待确认”。",
        required=True,
    )

    add_section_heading(doc, "企业主体信息", note="确认官网主体及母子公司关系")
    add_form_table(
        doc,
        [
            ("官网展示公司全称", "请填写工商登记全称"),
            ("公司简称", "请填写"),
            ("统一社会信用代码", "请填写"),
            ("成立时间", "____年__月__日"),
            ("法定代表人是否公开", "[ ] 公开，姓名：________  [ ] 不公开"),
            ("公司注册地址", "请填写"),
            ("官网对外展示地址", "请填写；如与注册地址不同请说明"),
            ("与北京首版科技有限公司的关系", "请准确描述母子公司、控股或关联关系"),
            ("母公司完整名称", "请填写"),
            ("可公开使用的关系表述", "示例仅供确认：北京首版科技有限公司为母公司"),
            ("资质归属划分", "请说明哪些资质属于首版认证，哪些属于母公司首版科技"),
        ],
        multiline_labels={"与北京首版科技有限公司的关系", "可公开使用的关系表述", "资质归属划分"},
    )

    add_section_heading(doc, "联系方式", note="用于联系页、页脚与联系表单")
    add_form_table(
        doc,
        [
            ("对外联系电话", "请填写"),
            ("商务咨询邮箱", "请填写"),
            ("客服 / 投诉邮箱", "请填写；如与商务邮箱相同请注明"),
            ("办公地址", "请填写"),
            ("工作时间", "示例：工作日 9:00–18:00"),
            ("联系表单接收邮箱", "只填写接收邮箱，不填写密码、授权码或密钥"),
            ("微信 / 公众号 / 二维码", "[ ] 不公开  [ ] 公开，名称或文件：________"),
            ("地图定位", "[ ] 不需要  [ ] 需要，地图地址：________"),
        ],
    )

    add_page_break(doc)
    add_section_heading(doc, "备案与法律信息", note="影响页脚、隐私政策及数据处理说明")
    add_form_table(
        doc,
        [
            ("ICP备案号", "请填写"),
            ("ICP备案链接", "默认：https://beian.miit.gov.cn/；如有指定链接请填写"),
            ("公安联网备案号", "请填写；如暂无请注明"),
            ("公安备案链接", "请填写"),
            ("正式域名", "请填写最终生产域名"),
            ("隐私政策生效日期", "____年__月__日"),
            ("联系表单数据保留期限", "请填写，例如：处理完成后 180 天"),
            ("个人信息投诉 / 删除联系人", "请填写邮箱或电话"),
            ("网站统计工具", "[ ] 不使用  [ ] 百度统计  [ ] 其他：________"),
            ("是否存在数据跨境", "[ ] 否  [ ] 是，说明：________"),
            ("法律内容审核人", "姓名 / 职务 / 联系方式"),
        ],
        multiline_labels={"联系表单数据保留期限", "是否存在数据跨境"},
    )

    add_section_heading(doc, "公司定位与核心文案", note="统一首页、关于页和搜索摘要口径")
    add_form_table(
        doc,
        [
            ("一句话定位", "当前文案供确认：AIGC时代数字人格权资产基础设施"),
            ("100–150字公司简介", "请填写可公开的正式版本"),
            ("最重要的三项能力", "1. ________  2. ________  3. ________"),
            ("主要服务对象", "[ ] 艺人  [ ] 创作者  [ ] 经纪机构  [ ] 品牌方  [ ] 平台企业  [ ] AI公司  [ ] 其他"),
            ("与普通版权登记 / 存证平台的区别", "请填写差异化价值，不使用无法验证的绝对化表述"),
            ("希望访客记住的一句话", "请填写"),
            ("禁止使用的表述", "列出敏感、夸大、未经授权或不应公开的词语与承诺"),
        ],
        multiline_labels={"100–150字公司简介", "与普通版权登记 / 存证平台的区别", "禁止使用的表述"},
    )

    add_page_break(doc)
    add_section_heading(doc, "服务与产品信息", note="先确认服务范围，再补齐每项服务详情")
    add_matrix(
        doc,
        ["服务", "是否提供", "服务成果 / 交付物", "常规周期", "价格展示"],
        [
            ["数字人格权资产认证", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["声音与声纹资产服务", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["肖像与影像资产服务", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["AIGC侵权监测", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["证据固定与维权协同", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["商业授权与合规使用", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
            ["其他服务：________", "[ ] 是  [ ] 否", "请填写", "请填写", "[ ] 公开  [ ] 咨询"],
        ],
        [2400, 1260, 3000, 1200, 1500],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_kicker(doc, "单项服务详情（每项服务复制填写一份）", color=COBALT_DARK, after=5)
    add_form_table(
        doc,
        [
            ("服务名称", "请填写"),
            ("服务对象", "请填写"),
            ("解决的问题", "请填写"),
            ("服务流程", "请按顺序填写"),
            ("最终交付物", "请填写客户最终获得的材料、报告或记录"),
            ("常规周期", "请填写"),
            ("客户需准备的材料", "请填写"),
            ("能力边界 / 免责声明", "请写明不承诺内容及需要第三方配合的事项"),
            ("咨询去向", "请填写页面、表单、电话或邮箱"),
        ],
        multiline_labels={"解决的问题", "服务流程", "最终交付物", "客户需准备的材料", "能力边界 / 免责声明"},
    )

    add_page_break(doc)
    add_section_heading(doc, "星眸 AIPR", note="确认正式名称、状态和可公开能力")
    add_form_table(
        doc,
        [
            ("正式产品名称", "请填写"),
            ("AIPR 全称及含义", "请填写"),
            ("当前状态", "[ ] 已上线  [ ] 内测  [ ] 开发中  [ ] 能力集合"),
            ("独立系统 / 登录地址", "请填写；如不公开请注明"),
            ("真实界面截图", "请提供文件名或素材目录"),
            ("软件著作权 / 专利 / 技术证明", "请列出可公开的证明材料"),
            ("不能对外承诺的能力", "请填写"),
        ],
        multiline_labels={"软件著作权 / 专利 / 技术证明", "不能对外承诺的能力"},
    )
    add_matrix(
        doc,
        ["能力模块", "当前状态", "是否已实际使用", "是否可对外公开", "备注"],
        [
            ["数字人格权资产库", "请填写", "[ ] 是  [ ] 否", "[ ] 是  [ ] 否", "请填写"],
            ["AIGC侵权监测系统", "请填写", "[ ] 是  [ ] 否", "[ ] 是  [ ] 否", "请填写"],
            ["证据包与维权管理系统", "请填写", "[ ] 是  [ ] 否", "[ ] 是  [ ] 否", "请填写"],
            ["数字人格资产可信训练舱", "请填写", "[ ] 是  [ ] 否", "[ ] 是  [ ] 否", "请填写"],
            ["授权项目管理系统", "请填写", "[ ] 是  [ ] 否", "[ ] 是  [ ] 否", "请填写"],
        ],
        [2460, 1500, 1800, 1800, 1800],
    )

    add_page_break(doc)
    add_section_heading(doc, "资质与证书", note="必须核对持证主体与公开范围")
    add_matrix(
        doc,
        ["证书 / 资质", "持证主体", "证书编号", "颁发及有效期", "完整图片公开"],
        [
            ["国家高新技术企业", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["中关村高新技术企业", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["ICP / EDI 许可证", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["全国 SP 许可证", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["版权工作站", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["行业协会 / 理事单位", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
            ["其他证书：________", "请填写", "请填写", "颁发单位：____\n日期 / 有效期：____", "[ ] 是  [ ] 否"],
        ],
        [2200, 1700, 1600, 2460, 1400],
    )
    add_form_table(
        doc,
        [
            ("证书脱敏要求", "请说明需遮盖的姓名、编号、二维码、地址等信息"),
            ("资质归属表述", "“母公司资质与行业认可”是否准确：[ ] 是  [ ] 否，修订：________"),
            ("证书点击放大", "[ ] 允许  [ ] 不允许  [ ] 仅展示脱敏版"),
        ],
        multiline_labels={"证书脱敏要求", "资质归属表述"},
    )

    add_page_break(doc)
    add_section_heading(doc, "合作伙伴", note="Logo必须有真实合作关系与展示授权")
    add_matrix(
        doc,
        ["合作伙伴", "关系真实", "Logo展示授权", "是否加链接", "链接 / 备注"],
        [
            [name, "[ ] 是  [ ] 否", "[ ] 有  [ ] 无  [ ] 待确认", "[ ] 是  [ ] 否", "请填写"]
            for name in [
                "百度网盘",
                "百度文库",
                "第一视频",
                "Wemade",
                "视觉中国",
                "方正集团",
                "腾讯视频",
                "京东",
                "快手",
                "新片场",
                "学科网",
                "优酷",
            ]
        ],
        [1600, 1350, 2200, 1500, 2710],
    )
    add_form_table(
        doc,
        [
            ("需要删除的伙伴", "请填写"),
            ("需要新增的伙伴", "请填写"),
            ("Logo展示顺序", "请填写；如无要求则按现有顺序"),
            ("栏目名称", "“合作伙伴”是否准确：[ ] 是  [ ] 否，建议名称：________"),
        ],
    )

    add_page_break(doc)
    add_phase_banner(
        doc,
        "第二阶段：建议补充",
        "以下资料不会阻塞基础上线，但会显著提升官网的可信度、搜索表现和内容完整度。",
        required=False,
    )

    add_section_heading(doc, "新闻与研究内容", required=False, note="建议首发至少准备 3–5 篇")
    add_form_table(
        doc,
        [
            ("文章标题", "请填写"),
            ("内容类型", "[ ] 公司新闻  [ ] 行业研究  [ ] 合规分析  [ ] 活动动态"),
            ("发布日期", "____年__月__日"),
            ("作者 / 发布方", "请填写"),
            ("50–100字摘要", "请填写"),
            ("正文", "请提供文档文件名或链接"),
            ("封面图", "请提供文件名"),
            ("来源链接", "请填写；原创可留空"),
            ("是否允许全文公开", "[ ] 是  [ ] 否  [ ] 仅摘要"),
            ("关联服务", "请填写"),
        ],
        multiline_labels={"50–100字摘要"},
    )
    add_matrix(
        doc,
        ["建议首发专题", "资料状态", "负责人", "备注"],
        [
            ["公司与产品介绍", "[ ] 待准备  [ ] 已完成", "请填写", "请填写"],
            ["数字人格权资产研究", "[ ] 待准备  [ ] 已完成", "请填写", "请填写"],
            ["AIGC侵权监测与证据组织", "[ ] 待准备  [ ] 已完成", "请填写", "请填写"],
            ["声音、声纹、肖像、数字人授权边界", "[ ] 待准备  [ ] 已完成", "请填写", "请填写"],
            ["资质与真实行业活动动态", "[ ] 待准备  [ ] 已完成", "请填写", "请填写"],
        ],
        [3500, 2360, 1500, 2000],
    )

    add_page_break(doc)
    add_section_heading(doc, "公司发展与关于我们", required=False)
    add_form_table(
        doc,
        [
            ("成立背景", "请填写"),
            ("发展历程", "按“年份 + 事件”提供，可另附文档"),
            ("使命", "请填写"),
            ("愿景", "请填写"),
            ("价值观", "请填写"),
            ("办公 / 团队规模是否公开", "[ ] 公开  [ ] 不公开；可公开口径：________"),
            ("管理层 / 专家顾问", "请填写可公开姓名、职务与简介"),
            ("团队 / 活动照片", "请提供文件名及肖像授权说明"),
            ("分支机构 / 服务点", "请填写；如无可留空"),
        ],
        multiline_labels={"成立背景", "发展历程", "管理层 / 专家顾问"},
    )

    add_page_break(doc)
    add_section_heading(doc, "真实案例与数据", required=False, note="所有案例、评价和数据必须可核验")
    add_form_table(
        doc,
        [
            ("是否有可公开案例", "[ ] 有  [ ] 暂无"),
            ("客户名称是否公开", "[ ] 公开  [ ] 匿名"),
            ("客户Logo授权", "[ ] 有  [ ] 无  [ ] 待确认"),
            ("案例背景", "请填写"),
            ("提供的服务", "请填写"),
            ("可公开的结果", "请填写，不使用无法核验的数据"),
            ("客户评价与授权", "请提供原文及书面授权状态"),
            ("可核验数据", "请列出统计口径、时间范围和来源"),
            ("禁止公开的数据", "请填写"),
        ],
        multiline_labels={"案例背景", "提供的服务", "可公开的结果", "客户评价与授权", "可核验数据", "禁止公开的数据"},
    )

    add_page_break(doc)
    add_section_heading(doc, "图片与素材", required=False)
    add_matrix(
        doc,
        ["素材", "是否已有", "文件位置 / 文件名", "版权与使用限制"],
        [
            ["最终横向 Logo", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["深色 / 浅色 / 透明 Logo", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["办公园区实景照片", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["会议 / 活动 / 团队照片", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["星眸 AIPR 真实截图", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["证书与资质原图", "[ ] 有  [ ] 无", "请填写", "请填写"],
            ["新闻与研究封面", "[ ] 有  [ ] 无", "请填写", "请填写"],
        ],
        [2400, 1500, 3160, 2300],
    )
    add_form_table(
        doc,
        [
            ("统一脱敏要求", "请填写"),
            ("版权 / 肖像权限制", "请填写"),
            ("图片质量要求确认", "[ ] 原图  [ ] 宽度不低于1600px  [ ] 避免聊天软件重复压缩"),
        ],
        multiline_labels={"统一脱敏要求", "版权 / 肖像权限制"},
    )

    add_page_break(doc)
    add_section_heading(doc, "SEO 与对外分享", required=False)
    add_form_table(
        doc,
        [
            ("官网正式名称", "请填写"),
            ("首页浏览器标题", "建议供确认：北京首版认证有限公司｜数字人格权资产服务"),
            ("搜索摘要（80–120字）", "请填写"),
            ("核心关键词", "请填写 5–10 个"),
            ("微信分享标题", "请填写"),
            ("微信分享摘要", "请填写"),
            ("微信分享封面", "请提供文件名"),
            ("百度站长平台验证", "[ ] 需要  [ ] 不需要  [ ] 待确认"),
            ("是否允许搜索引擎收录", "[ ] 允许  [ ] 暂不允许"),
        ],
        multiline_labels={"搜索摘要（80–120字）"},
    )

    add_section_heading(doc, "上线安排", required=False, note="用于最终发布与验收")
    add_form_table(
        doc,
        [
            ("目标上线日期", "____年__月__日"),
            ("正式域名", "请填写"),
            ("域名负责人", "姓名 / 联系方式"),
            ("服务器 / 部署平台", "请填写"),
            ("网站备案状态", "[ ] 已完成  [ ] 进行中  [ ] 未开始"),
            ("生产邮件服务", "请填写服务商与接收邮箱，不填写密钥"),
            ("内容审核人", "姓名 / 职务"),
            ("法律与资质审核人", "姓名 / 职务"),
            ("最终批准人", "姓名 / 职务"),
            ("是否需要预览地址", "[ ] 需要  [ ] 不需要"),
        ],
    )

    add_notes_box(doc, "最终上线备注 / 未解决事项")

    add_page_break(doc)
    add_kicker(doc, "提交前检查", color=COBALT, after=3)
    closing = doc.add_paragraph()
    closing.paragraph_format.space_after = Pt(12)
    run = closing.add_run("资料交付确认")
    set_run_font(run, size=22, color=INK, bold=True)
    add_matrix(
        doc,
        ["检查项", "负责人", "状态", "确认日期"],
        [
            ["企业主体与母子公司关系已核对", "请填写", "[ ] 已确认", "____年__月__日"],
            ["联系方式与表单接收邮箱已测试", "请填写", "[ ] 已确认", "____年__月__日"],
            ["ICP备案、隐私政策及法律信息已审核", "请填写", "[ ] 已确认", "____年__月__日"],
            ["服务范围、交付物与能力边界已审核", "请填写", "[ ] 已确认", "____年__月__日"],
            ["资质持证主体、编号和公开范围已核对", "请填写", "[ ] 已确认", "____年__月__日"],
            ["合作伙伴关系与Logo展示授权已确认", "请填写", "[ ] 已确认", "____年__月__日"],
            ["全部图片已确认版权、肖像权和脱敏要求", "请填写", "[ ] 已确认", "____年__月__日"],
            ["桌面端与移动端最终页面已验收", "请填写", "[ ] 已确认", "____年__月__日"],
        ],
        [4300, 1600, 1600, 1860],
    )
    add_form_table(
        doc,
        [
            ("最终批准人", "签字 / 姓名：________________"),
            ("批准日期", "____年__月__日"),
            ("上线结论", "[ ] 同意上线  [ ] 修改后上线  [ ] 暂缓上线"),
        ],
    )

    # Set updateFields so Word refreshes page numbers on open.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build website launch requirements DOCX.")
    parser.add_argument("output", type=Path, help="Output DOCX path")
    args = parser.parse_args()
    build_document(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
